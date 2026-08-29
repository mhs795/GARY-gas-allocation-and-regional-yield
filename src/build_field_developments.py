"""
Extract AEMO's 2026 GSOO field development pipeline into a tidy CSV.

Source workbook (data/other data/2026 GSOO/2026-gsoo-supply-data/):
  - G26 Processing Transmission Storage Facilities.xlsx, sheet "Field Developments"

WHY THIS EXISTS. GARY's supply side had no BACKFILL: every basin was a single row
that declined on a curve, and the only way new gas entered was the handful of
Type=Terminal candidates in expansion_options.csv. AEMO's own Figure 27 says that
is not what the south looks like -- existing southern fields do collapse (304 PJ/yr
in 2025 to 5 PJ/yr by 2044) but committed and anticipated DEVELOPMENTS backfill
them to ~280 PJ/yr. Without those developments a hard reserves cutoff models the
collapse and not the replacement, which is why the first attempt at one produced
6,419 TJ of shortage rather than a market (see model._declined_capacity).

This is the candidate list that backfill has to be built from. AEMO publishes, per
development: a Status, the basin, a first-production year and -- for some -- a
deliverability in TJ/day.

WHAT AEMO DOES NOT PUBLISH, and so is NOT in this file:
  * CapEx. No development capital appears anywhere in the GSOO supply data. A
    candidate built from this file has to carry GARY's own CapEx, flagged as such
    in expansion_options.csv exactly like the existing GARY-costed rows.
  * A deliverability for many rows. The Otway pipeline (Annie, Juliet, Nestor,
    Elanora, Wobbegong) has first-production YEARS but "N/A" capacity, and the
    Amadeus fields are "Commercial in confidence". Those rows are emitted with a
    blank Capacity_TJd rather than a guess.

Also emits AEMO's Figure 27 southern supply envelope, which is the aggregate the
individual developments have to add up to -- the calibration target for any
backfill capacity GARY has to invent where AEMO publishes none.

Outputs (data/gsoo/):
  - field_developments.csv     Name, Status, Basin, FirstProduction, Capacity_TJd,
                               GaryNode, IsCandidate, Note
  - southern_supply_envelope.csv  Year, Category, PJ_per_year   (2026 GSOO columns)
"""
import os
import re

import openpyxl
import pandas as pd

import params as P

BASE = os.path.dirname(__file__)
DATA = os.path.join(BASE, "data")
SRC = os.path.join(DATA, "other data", "2026 GSOO", "2026-gsoo-supply-data")
FAC = os.path.join(SRC, "G26 Processing Transmission Storage Facilities.xlsx")
FIG = os.path.join(SRC, "2026-gsoo-report-figures-and-data.xlsx")
OUT = os.path.join(DATA, "gsoo")

HORIZON_START = P.get_int('horizon_start', 2025)

# AEMO's basin names -> the GARY supply node that basin's gas enters the network
# at. A basin GARY does not carry (Galilee, Gunnedah, Bass) maps to None and its
# developments are extracted but never become candidates.
BASIN_TO_NODE = {
    'surat': 'Surat', 'bowen': 'Surat', 'surat-bowen': 'Surat',
    'surat/bowen': 'Surat', 'north bowen': 'Surat',
    'gippsland': 'Gippsland', 'offshore gippsland': 'Gippsland',
    'otway': 'Iona',
    'cooper': 'Moomba', 'eromanga': 'Moomba', 'cooper eromanga': 'Moomba',
    'amadeus': 'Amadeus',
    'beetaloo': 'Beetaloo',
}

# A Status word that means the field is ALREADY in GARY's base capacity. Those rows
# are extracted for completeness but are never candidates -- counting them again
# would double the deliverability the basin decline curves already carry. Matched
# anywhere in the cell, not just at the start, because AEMO writes compound
# statuses like "Sole - online; Manta - pre-FID".
ONLINE = ('online', 'existing', 'producing', 'production')

# AEMO's own header, spelled out in full so a change of wording in a later GSOO
# fails loudly here rather than silently emitting an empty candidate set.
CAP_COL = 'Production Capacity (TJ per day unless otherwise specified)'


def _norm_basin(text):
    """AEMO's free-text Basin cell -> a GARY node, or None."""
    t = str(text or '').lower()
    t = re.sub(r'\bbasin\b', ' ', t)
    for key in sorted(BASIN_TO_NODE, key=len, reverse=True):
        if key in t:
            return BASIN_TO_NODE[key]
    return None


def _capacity(text):
    """First plain TJ/day figure in AEMO's capacity cell, or None.

    The cell is free text: "375", "1486-1539", "40 (Pilot) 1000 (Full field)",
    "Not Currently Available", "Commercial in confidence". A RANGE takes its LOW
    end and a pilot/full pair takes the pilot, so the extract never overstates
    deliverability; the raw string is kept in Note so the choice stays visible.
    """
    s = str(text or '').strip()
    if not s or s.lower().startswith(('n/a', 'not ', 'commercial', 'refer', 'unknown')):
        return None
    m = re.search(r'(\d[\d,]*\.?\d*)', s.replace(' ', ''))
    return float(m.group(1).replace(',', '')) if m else None


def _first_year(text):
    """First 4-digit year in AEMO's first-production cell, or None."""
    m = re.search(r'(20\d{2}|19\d{2})', str(text or ''))
    return int(m.group(1)) if m else None


def developments():
    """The Field Developments sheet as a tidy frame."""
    ws = openpyxl.load_workbook(FAC, data_only=True)['Field Developments']
    hdr, rows = None, []
    for r in ws.iter_rows(values_only=True):
        v = [('' if c is None else str(c).replace('\n', ' ').strip()) for c in r]
        if len(v) > 1 and v[1] == 'Name':
            hdr = v
            continue
        if hdr and len(v) > 1 and v[1]:
            rows.append(v)
    idx = {name: i for i, name in enumerate(hdr) if name}
    missing = [c for c in ('Name', 'Status', 'Basin', 'First Production', CAP_COL)
               if c not in idx]
    if missing:
        raise KeyError(f"{FAC}: 'Field Developments' is missing {missing}. AEMO "
                       f"renamed a column; update the constants in this module.")
    out = []
    for v in rows:
        def col(name):
            i = idx.get(name)
            return v[i] if i is not None and i < len(v) else ''
        status = col('Status')
        node = _norm_basin(col('Basin'))
        cap = _capacity(col(CAP_COL))
        first = _first_year(col('First Production'))
        online = any(w in status.lower() for w in ONLINE)
        out.append({
            'Name': col('Name'),
            'Status': status,
            'Basin': col('Basin'),
            'GaryNode': node or '',
            'FirstProduction': first or '',
            'Capacity_TJd': cap if cap is not None else '',
            # A CANDIDATE is a development GARY could build on top of what it
            # already models. Four things all have to hold, and each exclusion is
            # there to stop a specific kind of double count:
            #   * the basin is one GARY carries              (else nowhere to put it)
            #   * the status says it is not already producing (else the basin decline
            #     curve already carries this deliverability)
            #   * AEMO published a positive deliverability    (else GARY would be
            #     inventing the capacity, not sourcing it)
            #   * first production falls inside the horizon   (a field flowing since
            #     2009 or 2019 is in the base capacity too, whatever its status says)
            'IsCandidate': (bool(node) and not online and cap is not None
                            and cap > 0 and first is not None and first >= HORIZON_START),
            # Recorded but NOT a candidate: AEMO names the development and gives it
            # a start year, but publishes no deliverability. The Otway pipeline
            # (Annie, Juliet, Nestor, Elanora, Wobbegong) is all of this, and it is
            # the backfill for the one basin that collapsed hardest under a stock
            # cutoff -- so it is flagged rather than silently dropped.
            'CapacityUnpublished': (bool(node) and not online and cap is None
                                    and first is not None and first >= HORIZON_START),
            'Note': ' | '.join(x for x in (col(CAP_COL),
                                           col('First Production')) if x)[:200],
        })
    return pd.DataFrame(out)


def southern_envelope():
    """AEMO Figure 27: southern production by category, 2026 GSOO columns only."""
    ws = openpyxl.load_workbook(FIG, data_only=True)['Figure 27']
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    header = next(r for r in grid if any(isinstance(c, str) and 'Existing and Com' in c
                                         for c in r if c is not None))
    cols = {i: str(c).strip() for i, c in enumerate(header)
            if isinstance(c, str) and c.strip().startswith('2026 GSOO')}
    out, year = [], None
    for r in grid:
        for c in r:
            if isinstance(c, (int, float)) and 2020 <= c <= 2060 and float(c).is_integer():
                year = int(c)
            elif isinstance(c, str) and re.fullmatch(r'20\d{2}', c.strip()):
                year = int(c)
        if year is None:
            continue
        for i, label in cols.items():
            v = r[i] if i < len(r) else None
            if isinstance(v, (int, float)):
                out.append({'Year': year,
                            'Category': label.replace('2026 GSOO ', ''),
                            'PJ_per_year': round(float(v), 3)})
    return pd.DataFrame(out).drop_duplicates(['Year', 'Category'])


def derived_capex():
    """Development capital implied by AEMO's own numbers, per basin and per project.

    AEMO publishes no CapEx, only a blended $/GJ that its Production Costs note says
    includes "operating cost, capital costs, royalty, tax and a return on capital".
    supply.csv splits that: a 2C row carries the basin's OPERATING basis in `Cost`
    and AEMO's published full cost in `AEMOFullCost`. The capital is the gap, over
    the resource that capital develops, shared across the basin's developments pro
    rata on the deliverability each brings.

    This is the check that expansion_options.csv has not drifted from supply.csv.
    """
    sup = pd.read_csv(os.path.join(DATA, 'supply.csv'))
    exp = pd.read_csv(os.path.join(DATA, 'expansion_options.csv'))
    c2c = sup[(sup['IsPotential']) & (sup['Tranche'] == '2C')]
    per_basin = {r['Node']: (r['AEMOFullCost'] - r['Cost']) * r['Reserves_PJ'] * 1e6
                 for _, r in c2c.iterrows()}
    out = []
    for node, capex in per_basin.items():
        devs = exp[(exp['Type'] == 'Terminal') & (exp['Target'] == node)]
        total = devs['NewCapacity'].sum()
        for _, d in devs.iterrows():
            want = capex * d['NewCapacity'] / total
            out.append({'Name': d['Name'], 'Basin': node,
                        'Expected_CapEx': round(want),
                        'In_File': int(d['CapEx']),
                        'Matches': abs(want - d['CapEx']) <= max(1.0, 0.001 * want)})
    return pd.DataFrame(out), per_basin


def main():
    os.makedirs(OUT, exist_ok=True)
    dev = developments()
    dev.to_csv(os.path.join(OUT, 'field_developments.csv'), index=False)
    env = southern_envelope()
    env.to_csv(os.path.join(OUT, 'southern_supply_envelope.csv'), index=False)
    cand = dev[dev['IsCandidate']]
    print(f"field_developments.csv: {len(dev)} developments, "
          f"{len(cand)} usable as candidates")
    for _, r in cand.iterrows():
        print(f"   {r['GaryNode']:<10} {r['Name'][:34]:<34} "
              f"{r['Capacity_TJd']:>7} TJ/d  {r['FirstProduction'] or '?'}  [{r['Status'][:26]}]")
    unpub = dev[dev['CapacityUnpublished']]
    if len(unpub):
        print(f"   -- {len(unpub)} more named with a start year but NO published "
              f"deliverability:")
        for _, r in unpub.iterrows():
            print(f"   {r['GaryNode']:<10} {r['Name'][:34]:<34} "
                  f"{'?':>7}       {r['FirstProduction']}  [{r['Status'][:26]}]")
    print(f"southern_supply_envelope.csv: {len(env)} rows, "
          f"{env['Year'].min()}-{env['Year'].max()}")
    return dev, env


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--capex', action='store_true',
                    help="Recompute field-development CapEx from supply.csv and "
                         "check expansion_options.csv still agrees with it")
    args = ap.parse_args()
    if args.capex:
        df, per_basin = derived_capex()
        print("Development capital implied by AEMO's 2C cost gap:")
        for node, v in sorted(per_basin.items(), key=lambda kv: -kv[1]):
            print(f"   {node:<12} ${v/1e9:>6.2f}bn")
        print()
        df['Expected_$bn'] = (df['Expected_CapEx'] / 1e9).round(2)
        df['InFile_$bn'] = (df['In_File'] / 1e9).round(2)
        print(df[['Name', 'Basin', 'Expected_$bn', 'InFile_$bn', 'Matches']]
              .to_string(index=False))
        bad = df[~df['Matches']]
        print(f"\n{'DRIFT: ' + str(len(bad)) + ' row(s) disagree' if len(bad) else 'All rows agree with supply.csv.'}")
    else:
        main()
